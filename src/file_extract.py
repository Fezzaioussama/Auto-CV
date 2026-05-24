"""Extract plain text from uploaded CV files (PDF / DOCX / TXT / TEX).

Most job seekers do not have a LaTeX CV — they have a PDF or Word document. This
module turns those uploads into text the rest of the pipeline can work with.
Heavy parsers are imported lazily so a missing optional dependency degrades to a
clear error instead of breaking app startup.
"""

from __future__ import annotations

import io
from typing import Tuple


class ExtractionError(Exception):
    """Raised when a file cannot be read or is unsupported."""


def _ext(filename: str) -> str:
    return (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise ExtractionError(
            "PDF support is not installed (pip install pdfplumber)."
        ) from exc

    chunks = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                chunks.append(page.extract_text() or "")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not read the PDF file.") from exc
    return "\n".join(chunks).strip()


def extract_text_from_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise ExtractionError(
            "Word support is not installed (pip install python-docx)."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not read the Word document.") from exc

    lines = [p.text for p in document.paragraphs]
    # Pull text out of tables too (skills grids, etc.).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(line for line in lines if line is not None).strip()


def extract_cv_text(filename: str, data: bytes) -> Tuple[str, str]:
    """Return ``(text, source_format)`` for an uploaded file.

    ``source_format`` is one of ``pdf`` | ``docx`` | ``tex`` | ``text``.
    Raises :class:`ExtractionError` for unsupported types or empty results.
    """
    ext = _ext(filename)
    if ext == "pdf":
        text, fmt = extract_text_from_pdf(data), "pdf"
    elif ext in ("docx", "doc"):
        text, fmt = extract_text_from_docx(data), "docx"
    elif ext == "tex":
        text, fmt = data.decode("utf-8", errors="replace"), "tex"
    elif ext in ("txt", ""):
        text, fmt = data.decode("utf-8", errors="replace"), "text"
    else:
        raise ExtractionError(f"Unsupported file type: .{ext}")

    if not text or not text.strip():
        raise ExtractionError(
            "No text could be extracted (the file may be scanned/image-only)."
        )
    return text.strip(), fmt
