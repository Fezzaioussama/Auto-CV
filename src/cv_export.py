"""Export a CV to recruiter-friendly formats: Word (.docx) and plain text.

The optimizer's working format is LaTeX, but most recruiters and job portals want
a Word document or plain text they can paste into an ATS form. This module turns
the (tailored) LaTeX into a clean, readable .docx / .txt — it is intentionally a
pragmatic converter, not a full TeX engine: it keeps the structure (headings,
bullets, paragraphs) and the text, and drops formatting commands.
"""

from __future__ import annotations

import io
import re
from typing import List, Dict, Tuple


class ExportError(Exception):
    """Raised when a CV cannot be exported in the requested format."""


# ---------------------------------------------------------------------------
# LaTeX → structured blocks
# ---------------------------------------------------------------------------

_SECTION_RE = re.compile(r"\\(?:section)\*?\{(.*?)\}")
_SUBSECTION_RE = re.compile(r"\\(?:subsection|subsubsection)\*?\{(.*?)\}")
_ITEM_RE = re.compile(r"^\s*\\item\b\s*(.*)$")
_HREF_RE = re.compile(r"\\href\{([^}]*)\}\{([^}]*)\}")
# Formatting wrappers whose argument we keep verbatim.
_WRAP_RE = re.compile(
    r"\\(?:textbf|textit|emph|underline|texttt|textsc|textnormal|mbox|text|"
    r"large|Large|LARGE|small|footnotesize|textsf|textrm)\{([^{}]*)\}"
)
# Any remaining "\cmd[opt]{arg}" — keep the argument, drop the command.
_CMD_ARG_RE = re.compile(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?\{([^{}]*)\}")
# Bare "\cmd" or "\cmd[opt]" with no argument — drop entirely.
_CMD_RE = re.compile(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?")
_ENV_RE = re.compile(r"\\(?:begin|end)\{[^}]*\}")
_COMMENT_RE = re.compile(r"(?<!\\)%.*$")


def _clean_inline(text: str) -> str:
    """Reduce a LaTeX fragment to readable plain text."""
    if not text:
        return ""

    def _href(m):
        url = m.group(1).replace("mailto:", "")
        label = m.group(2)
        if not label or label == url:
            return url
        return f"{label} ({url})"

    text = _HREF_RE.sub(_href, text)
    # Resolve nested formatting wrappers a few times.
    for _ in range(4):
        new = _WRAP_RE.sub(r"\1", text)
        if new == text:
            break
        text = new
    text = _ENV_RE.sub("", text)
    for _ in range(3):
        new = _CMD_ARG_RE.sub(r"\1", text)
        if new == text:
            break
        text = new
    text = _CMD_RE.sub("", text)
    text = text.replace("{", "").replace("}", "")
    # Unescape LaTeX specials and tidy spacing characters.
    text = re.sub(r"\\([&%$#_{}])", r"\1", text)
    text = text.replace("~", " ").replace("\\,", " ").replace("\\&", "&")
    text = re.sub(r"[ \t]+", " ", text).strip()
    return text


def _strip_to_body(latex: str) -> str:
    """Return the document body (between begin/end document) when present."""
    begin = latex.find(r"\begin{document}")
    if begin != -1:
        latex = latex[begin + len(r"\begin{document}"):]
    end = latex.find(r"\end{document}")
    if end != -1:
        latex = latex[:end]
    return latex


def latex_to_blocks(latex: str) -> List[Dict[str, str]]:
    """Parse CV LaTeX into ordered blocks: heading / subheading / bullet / text."""
    body = _strip_to_body(latex or "")
    body = _COMMENT_RE.sub("", body)
    body = body.replace(r"\\", "\n")  # LaTeX line breaks → real line breaks

    blocks: List[Dict[str, str]] = []
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue

        m = _SECTION_RE.search(line)
        if m and line.lstrip().startswith("\\section"):
            heading = _clean_inline(m.group(1))
            if heading:
                blocks.append({"kind": "heading", "text": heading})
            continue

        m = _SUBSECTION_RE.search(line)
        if m and line.lstrip().startswith("\\sub"):
            sub = _clean_inline(m.group(1))
            if sub:
                blocks.append({"kind": "subheading", "text": sub})
            continue

        m = _ITEM_RE.match(line)
        if m:
            text = _clean_inline(m.group(1))
            if text:
                blocks.append({"kind": "bullet", "text": text})
            continue

        if line.lstrip().startswith("\\begin") or line.lstrip().startswith("\\end"):
            continue

        text = _clean_inline(line)
        if text:
            blocks.append({"kind": "text", "text": text})
    return blocks


# ---------------------------------------------------------------------------
# Blocks → output formats
# ---------------------------------------------------------------------------


def blocks_to_text(blocks: List[Dict[str, str]]) -> str:
    out: List[str] = []
    for b in blocks:
        kind, text = b["kind"], b["text"]
        if kind == "heading":
            out.append("")
            out.append(text.upper())
            out.append("-" * len(text))
        elif kind == "subheading":
            out.append("")
            out.append(text)
        elif kind == "bullet":
            out.append(f"  • {text}")
        else:
            out.append(text)
    return "\n".join(out).strip() + "\n"


def blocks_to_docx(blocks: List[Dict[str, str]]) -> bytes:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise ExportError("Word export is not available (pip install python-docx).") from exc

    document = docx.Document()
    for b in blocks:
        kind, text = b["kind"], b["text"]
        if kind == "heading":
            document.add_heading(text, level=1)
        elif kind == "subheading":
            document.add_heading(text, level=2)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


_FORMATS = {
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "txt": ("text/plain; charset=utf-8", "txt"),
}


def export_cv(latex: str, fmt: str) -> Tuple[bytes, str, str]:
    """Return ``(bytes, mimetype, extension)`` for the requested format.

    ``fmt`` is ``docx`` or ``txt``. Raises :class:`ExportError` on bad input.
    """
    fmt = (fmt or "").lower().strip()
    if fmt not in _FORMATS:
        raise ExportError(f"Unsupported export format: {fmt!r}. Use 'docx' or 'txt'.")
    if not (latex or "").strip():
        raise ExportError("No CV content to export.")

    blocks = latex_to_blocks(latex)
    if not blocks:
        # Nothing parsed (e.g. already plain text) — fall back to raw text.
        blocks = [{"kind": "text", "text": line.strip()}
                  for line in (latex or "").splitlines() if line.strip()]

    mimetype, ext = _FORMATS[fmt]
    data = blocks_to_docx(blocks) if fmt == "docx" else blocks_to_text(blocks).encode("utf-8")
    return data, mimetype, ext
