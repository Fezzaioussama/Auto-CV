"""Extract plain text from uploaded CV files (PDF / DOCX / image / TXT / TEX).

Most job seekers do not have a LaTeX CV — they have a PDF, a Word document, or
even a photo/scan of a printed CV. This module turns those uploads into text the
rest of the pipeline can work with, picking the right strategy per file type:

* PDF / DOCX with a text layer → parse the text directly (fast, lossless).
* PDF with no text layer (a scan) or an image file → OCR via :mod:`ocr`
  (Tesseract locally, else a vision LLM).

Heavy parsers are imported lazily so a missing optional dependency degrades to a
clear error instead of breaking app startup.
"""

from __future__ import annotations

import io
from typing import Tuple


class ExtractionError(Exception):
    """Raised when a file cannot be read or is unsupported."""


# Image formats we accept and route through OCR.
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp", "gif"}

# Below this many non-whitespace characters, a PDF's text layer is treated as
# empty (i.e. the PDF is really a scan) and we fall back to OCR.
_SCANNED_PDF_TEXT_THRESHOLD = 40


def _ext(filename: str) -> str:
    return (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise ExtractionError(
            "PDF support is not installed (pip install pdfplumber)."
        ) from exc

    chunks = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                chunks.append(page.extract_text() or "")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not read the PDF file.") from exc
    return "\n".join(chunks).strip()


def extract_text_from_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise ExtractionError(
            "Word support is not installed (pip install python-docx)."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not read the Word document.") from exc

    lines = [p.text for p in document.paragraphs]
    # Pull text out of tables too (skills grids, etc.).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(line for line in lines if line is not None).strip()


def _ocr():
    """Lazy import of the OCR layer (keeps optional deps out of import time)."""
    try:  # importable both as a bare module (main.py) and as the src package
        import ocr
    except ImportError:  # pragma: no cover
        from . import ocr  # type: ignore
    return ocr


def extract_text_from_image(data: bytes) -> str:
    """OCR an uploaded image file (photo/screenshot/scan of a CV)."""
    ocr = _ocr()
    try:
        return ocr.ocr_image_bytes(data)
    except ocr.OCRError as exc:
        raise ExtractionError(str(exc)) from exc


def _extract_pdf_with_ocr_fallback(data: bytes) -> str:
    """PDF text layer first; OCR the rasterized pages if it looks scanned."""
    text = extract_text_from_pdf(data)
    if len(text.replace(" ", "").strip()) >= _SCANNED_PDF_TEXT_THRESHOLD:
        return text  # real text layer — no need to OCR

    # Sparse/empty text layer: this is almost certainly a scan. OCR it.
    ocr = _ocr()
    try:
        ocr_text = ocr.ocr_pdf_bytes(data)
    except ocr.OCRError as exc:
        if text.strip():
            return text  # keep whatever little we got rather than failing hard
        raise ExtractionError(str(exc)) from exc
    # Prefer whichever source produced more text.
    return ocr_text if len(ocr_text) >= len(text) else text


def extract_cv_text(filename: str, data: bytes) -> Tuple[str, str]:
    """Return ``(text, source_format)`` for an uploaded file.

    ``source_format`` is one of ``pdf`` | ``docx`` | ``image`` | ``tex`` |
    ``text``. PDFs without a text layer and image uploads are read with OCR.
    Raises :class:`ExtractionError` for unsupported types or empty results.
    """
    ext = _ext(filename)
    if ext == "pdf":
        text, fmt = _extract_pdf_with_ocr_fallback(data), "pdf"
    elif ext in ("docx", "doc"):
        text, fmt = extract_text_from_docx(data), "docx"
    elif ext in IMAGE_EXTENSIONS:
        text, fmt = extract_text_from_image(data), "image"
    elif ext == "tex":
        text, fmt = data.decode("utf-8", errors="replace"), "tex"
    elif ext in ("txt", ""):
        text, fmt = data.decode("utf-8", errors="replace"), "text"
    else:
        raise ExtractionError(f"Unsupported file type: .{ext}")

    if not text or not text.strip():
        raise ExtractionError(
            "No text could be extracted (the file may be scanned/image-only)."
        )
    return text.strip(), fmt
